from __future__ import annotations

import argparse
import json
import re
from datetime import date
from pathlib import Path

from docx import Document


def normalize_text(value: str) -> str:
    value = value.replace("\r", " ").replace("\n", " ").replace("\xa0", " ")
    return re.sub(r"\s+", " ", value).strip()


def sql_literal(value: str | None) -> str:
    if value is None:
        return "NULL"
    return "'" + value.replace("\\", "\\\\").replace("'", "''") + "'"


def read_codebook_rows(docx_path: Path) -> list[dict[str, str | int]]:
    document = Document(str(docx_path))
    if not document.tables:
        raise ValueError(f"Keine Tabelle in Word-Datei gefunden: {docx_path}")

    rows: list[dict[str, str | int]] = []
    table = document.tables[0]

    for row_index, row in enumerate(table.rows):
        cells = [normalize_text(cell.text) for cell in row.cells]
        while len(cells) < 5:
            cells.append("")

        if row_index == 0 and cells[0].lower() == "nr":
            continue

        if not any(cells):
            continue

        number_text, code_group, code, example_text, coding_rule = cells[:5]
        if not number_text:
            continue

        try:
            number = int(number_text)
        except ValueError as exc:
            raise ValueError(
                f"Ungueltige Nummer in Zeile {row_index + 1}: {number_text!r}"
            ) from exc

        rows.append(
            {
                "number": number,
                "codeGroup": code_group,
                "code": code,
                "exampleText": example_text,
                "codingRule": coding_rule,
            }
        )

    return rows


def build_codebook_payload(source_document: Path, labels: list[dict[str, str | int]]) -> dict:
    return {
        "project": "CSS",
        "sourceDocument": str(source_document),
        "description": "Import der CSS-Codeliste aus Word inklusive Nummer, Codegruppe, Code, Beispiel und Codierregel.",
        "version": str(date.today()),
        "labelCount": len(labels),
        "labels": labels,
    }


def build_sql(project_ids: list[int], labels: list[dict[str, str | int]], source_document: Path) -> str:
    # Match stable external numbers before labels: a renamed label must retain its
    # keyword ID so existing statement assignments remain valid.
    projects = ", ".join(str(project_id) for project_id in project_ids)
    lines = [
        "-- Generiert aus der Word-Codeliste",
        f"-- Quelle: {source_document}",
        f"-- Generiert am: {date.today().isoformat()}",
        f"-- CSS-Projekte: {projects}. Bestehende IDs und Codierungen bleiben erhalten.",
        "SET NAMES utf8mb4;",
        "DROP TEMPORARY TABLE IF EXISTS css_codebook_update;",
        "CREATE TEMPORARY TABLE css_codebook_update (",
        "  external_code_nr VARCHAR(16) NOT NULL PRIMARY KEY,",
        "  label VARCHAR(128) NOT NULL, codegruppe VARCHAR(256),",
        "  beispiel TEXT, codierregel TEXT",
        ") CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci;",
        "INSERT INTO css_codebook_update VALUES",
    ]
    lines.append(",\n".join(
        "  (" + ", ".join(sql_literal(str(label[key])) for key in
        ("number", "code", "codeGroup", "exampleText", "codingRule")) + ")"
        for label in labels
    ) + ";")
    lines += ["", "START TRANSACTION;"]
    for project_id in project_ids:
        lines += [
            f"-- Projekt {project_id}",
            "UPDATE ai_keywords k JOIN css_codebook_update c",
            "  ON k.external_code_nr = c.external_code_nr",
            "  OR ((k.external_code_nr IS NULL OR k.external_code_nr = '') AND k.label = c.label)",
            "SET k.external_code_nr = c.external_code_nr, k.label = c.label,",
            "    k.codegruppe = c.codegruppe, k.beispiel = c.beispiel, k.codierregel = c.codierregel",
            f"WHERE k.project_id = {project_id};",
            "INSERT INTO ai_keywords (project_id, external_code_nr, label, codegruppe, beispiel, codierregel, active, created_at)",
            f"SELECT {project_id}, c.external_code_nr, c.label, c.codegruppe, c.beispiel, c.codierregel, 1, NOW()",
            "FROM css_codebook_update c WHERE NOT EXISTS (",
            f"  SELECT 1 FROM ai_keywords k WHERE k.project_id = {project_id}",
            "    AND (k.external_code_nr = c.external_code_nr OR k.label = c.label)",
            ");",
            "",
        ]
    lines += [
        "COMMIT;",
        f"SELECT project_id, COUNT(*) AS code_count FROM ai_keywords WHERE project_id IN ({projects}) GROUP BY project_id;",
        "DROP TEMPORARY TABLE css_codebook_update;",
    ]
    return "\n".join(lines) + "\n"



def main() -> None:
    parser = argparse.ArgumentParser(
        description="Liest die CSS-Codeliste aus einer Word-Datei und erzeugt JSON + SQL fuer den Import."
    )
    parser.add_argument("--docx", required=True, help="Pfad zur Word-Datei mit der Codeliste.")
    parser.add_argument("--json-out", required=True, help="Zielpfad fuer die aktualisierte JSON-Codebook-Datei.")
    parser.add_argument("--sql-out", required=True, help="Zielpfad fuer das SQL-Importskript.")
    parser.add_argument(
        "--projects",
        nargs="+",
        type=int,
        default=[4, 5, 6, 7, 8],
        help="Projekt-IDs, in die die Codes per SQL importiert werden sollen.",
    )
    args = parser.parse_args()

    docx_path = Path(args.docx).expanduser().resolve()
    json_out = Path(args.json_out).expanduser().resolve()
    sql_out = Path(args.sql_out).expanduser().resolve()

    labels = read_codebook_rows(docx_path)
    payload = build_codebook_payload(docx_path, labels)
    sql_text = build_sql(args.projects, labels, docx_path)

    json_out.parent.mkdir(parents=True, exist_ok=True)
    sql_out.parent.mkdir(parents=True, exist_ok=True)

    json_out.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    sql_out.write_text(sql_text, encoding="utf-8")

    print(f"Codebook JSON geschrieben: {json_out}")
    print(f"SQL-Import geschrieben:   {sql_out}")
    print(f"Anzahl Codes:             {len(labels)}")
    print(f"Projekt-IDs:              {', '.join(str(project_id) for project_id in args.projects)}")


if __name__ == "__main__":
    main()
